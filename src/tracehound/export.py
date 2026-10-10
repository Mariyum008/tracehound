"""Turn a Markdown incident report into an editable Word document."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|_[^_]+_)")


def _add_inline(par, text: str):
    # Render **bold**, *italic* / _italic_ and `code` spans inside one paragraph
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def markdown_to_docx(markdown: str, title: str | None = None) -> bytes:
    """Convert the subset of Markdown our reports use (headings, lists, bold, code, rules) to .docx bytes."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    if title:
        doc.add_heading(title, level=0)
    in_code = False
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            run = doc.add_paragraph().add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(9.5)
            continue
        if not line.strip():
            continue
        if line.strip() in ("---", "***"):
            doc.add_paragraph().add_run("_" * 40).font.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            doc.add_heading(heading.group(2).replace("**", ""), level=min(len(heading.group(1)), 3))
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)", line)
        number = re.match(r"^(\s*)\d+[.)]\s+(.*)", line)
        if bullet or number:
            m = bullet or number
            style = "List Bullet" if bullet else "List Number"
            if len(m.group(1)) >= 2:
                style += " 2"
            _add_inline(doc.add_paragraph(style=style), m.group(2))
            continue
        _add_inline(doc.add_paragraph(), line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
